#!/usr/bin/env python3
"""Compara v06 y v07 celda por celda (valor, formula, estilo) y parte por parte del paquete.
Uso: verificar_v07.py <v06.xlsx> <v07.xlsx>
Salida: lista de diferencias; cualquier diferencia no listada en CAMBIOS-v07.md es un defecto.
"""
import sys
import zipfile
import openpyxl

DESPL = 12


def vol(ws):
    d = {}
    for row in ws.iter_rows():
        for c in row:
            est = (c.number_format, c.font.b, c.font.sz, c.font.name, c.alignment.horizontal, c.alignment.wrap_text,
                   c.border.left.style, c.border.right.style, c.border.top.style, c.border.bottom.style,
                   c.fill.fgColor.rgb if c.fill and c.fill.fill_type else None)
            d[c.coordinate] = (c.value, est)
    return d


def main(a, b):
    w6, w7 = openpyxl.load_workbook(a), openpyxl.load_workbook(b)
    assert w6.sheetnames == w7.sheetnames, (w6.sheetnames, w7.sheetnames)
    print("hojas:", w6.sheetnames)
    for n in w6.sheetnames:
        s6, s7 = w6[n], w7[n]
        d6, d7 = vol(s6), vol(s7)
        print("\n== hoja", repr(n))
        print("  celdas combinadas v06:", len(s6.merged_cells.ranges), " v07:", len(s7.merged_cells.ranges))
        print("  area de impresion v06:", s6.print_area, " v07:", s7.print_area)
        print("  imagenes v06:", len(s6._images), " v07:", len(s7._images))
        print("  dimension v06:", s6.dimensions, " v07:", s7.dimensions)
        print("  pageSetup scale/orient v06:", s6.page_setup.scale, s6.page_setup.orientation, "| v07:", s7.page_setup.scale, s7.page_setup.orientation)
        es_control = n == "CONTROL DE PAGOS"
        mapa_ok = set()
        for coord, (v, e) in d6.items():
            col = "".join(ch for ch in coord if ch.isalpha()); fila = int("".join(ch for ch in coord if ch.isdigit()))
            dest = "%s%d" % (col, fila + DESPL) if (es_control and fila >= 55) else coord
            mapa_ok.add(dest)
            v7, e7 = d7.get(dest, (None, None))
            if es_control and fila >= 55 and isinstance(v, str) and v.startswith("=SUM"):
                continue                       # formulas de totales: verificadas abajo
            if v != v7 or e != e7:
                print("  DIF %s -> %s: valor %r -> %r%s" % (coord, dest, v, v7, "" if e == e7 else "  ESTILO DISTINTO"))
        for coord in d7:
            if coord not in mapa_ok:
                col = "".join(ch for ch in coord if ch.isalpha()); fila = int("".join(ch for ch in coord if ch.isdigit()))
                if es_control and 55 <= fila <= 66:
                    continue                   # filas nuevas de cuotas
                print("  NUEVA", coord, d7[coord][0])
        if es_control:
            print("  C67:", d7["C67"][0], " E67:", d7["E67"][0])
            ok = all(d7["%s%d" % (c, r)][1] == d6["%s54" % c][1] for r in range(55, 67) for c in "ABCDE")
            print("  filas 55-66 con el mismo formato que la fila 54 (A-E):", ok)
            print("  A7:A66 vacias:", all(d7["A%d" % r][0] is None for r in range(7, 67)))

    z6, z7 = zipfile.ZipFile(a), zipfile.ZipFile(b)
    n6, n7 = {i.filename: z6.read(i.filename) for i in z6.infolist()}, {i.filename: z7.read(i.filename) for i in z7.infolist()}
    print("\n== partes del paquete")
    print("  partes solo en v06:", sorted(set(n6) - set(n7)), " solo en v07:", sorted(set(n7) - set(n6)))
    print("  identicas byte a byte:", sorted(k for k in n6 if k in n7 and n6[k] == n7[k]))
    print("  modificadas:", sorted(k for k in n6 if k in n7 and n6[k] != n7[k]))


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
