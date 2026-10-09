"""Verifica el contenido de los archivos devueltos por el SANDBOX. Llama a la Custom API (6, 30 y 60), decodifica y comprueba:
 (1) contra el esperado local (mismo codigo, net8.0): mismas partes del paquete y mismos bytes por parte;
 (2) con un oraculo independiente en Python (openpyxl + zipfile, sin el SDK): valores y tipos de celdas mapeadas, N cuotas,
     filas sobrantes limpias y con su estilo, totales, celdas combinadas, area de impresion, imagen, ausencia de TODAY();
 (3) LibreOffice -> PDF (paginas y texto)."""
import base64, io, zipfile, subprocess, datetime as dt, sys
from decimal import Decimal
import openpyxl
from sbx_lib import *

CONTROL_NEGATIVO = "--control-negativo" in sys.argv
dv = Dataverse(); e = leer_estado()
SAL = os.path.join(AQUI, "salida"); os.makedirs(SAL, exist_ok=True)
FORMATO, CONTROL = " FORMATO ARP ELECTRONICO", "CONTROL DE PAGOS"
plantilla_zip = zipfile.ZipFile(PLANTILLA)
wb_t = openpyxl.load_workbook(PLANTILLA)
fallos = []

def chequear(cond, msg):
    if not cond:
        fallos.append(msg)
    return cond

def esperado(plazo):
    """Oraculo independiente: escrito aparte del C# (Decimal de Python, fechas con calendario propio)."""
    def mas_meses(d, k):
        m = d.month - 1 + k
        return dt.datetime(d.year + m // 12, m % 12 + 1, d.day)
    inicio = dt.datetime(2026, 11, 8)
    cuota = Decimal("2150.35") if plazo == 60 else Decimal("4200.10")
    pagos = 12 if plazo >= 30 else 0
    cu = []
    for i in range(1, plazo + 1):
        f = mas_meses(inicio, i - 1)
        pagada = i <= pagos
        cu.append((i, f, cuota, f + dt.timedelta(days=1) if pagada else None, (cuota / 2 if i == 12 else cuota) if pagada else None))
    total = sum(c[2] for c in cu)
    pagado = sum(c[4] for c in cu if c[4] is not None)
    simples = {
        (FORMATO, "B7"): dt.datetime(2026, 10, 8), (FORMATO, "F7"): "Oficial de Prueba", (FORMATO, "B8"): "Ana Lucia Martinez Ruiz (ficticio)",
        (FORMATO, "F8"): "000-000000-0000X", (FORMATO, "B9"): "cliente.ficticio@ejemplo.com", (FORMATO, "F9"): "0000-0000",
        (FORMATO, "B12"): "0000-0000-0000-0000", (FORMATO, "D12"): "Temporal", (FORMATO, "F12"): 3, (FORMATO, "B13"): Decimal("125000.75"),
        (FORMATO, "D13"): Decimal("5000"), (FORMATO, "F13"): Decimal("120000.75"), (FORMATO, "B14"): plazo, (FORMATO, "D14"): cuota,
        (FORMATO, "F14"): total, (FORMATO, "B15"): inicio, (FORMATO, "D15"): mas_meses(inicio, plazo - 1), (FORMATO, "F15"): "0000000000000000",
        (FORMATO, "A18"): f"Datos ficticios del spike. Plazo de {plazo} cuotas.", (CONTROL, "B4"): "Ana Lucia Martinez Ruiz (ficticio)",
        (CONTROL, "E4"): dt.datetime(2026, 10, 8), (CONTROL, "C67"): total, (CONTROL, "E67"): pagado}
    return cu, simples

def igual(v, esp):
    if isinstance(esp, Decimal):
        return v is not None and not isinstance(v, str) and Decimal(str(v)) == esp
    if isinstance(esp, dt.datetime):
        return isinstance(v, dt.datetime) and v == esp
    return v == esp and type(v) is type(esp)

resumen = []
for plazo in (6, 30, 60):
    if CONTROL_NEGATIVO:   # control de sensibilidad: se le da al oraculo el archivo de OTRO plazo; tiene que fallar
        archivo = open(os.path.join(SAL, f"sandbox-{6 if plazo != 6 else 60}.xlsx"), "rb").read()
        b = {"TamanoBytes": len(archivo)}
    else:
        est, b, _ = dv.call("POST", API, {"VersionId": e["filas"]["v07 borrador (huella correcta)"], "Plazo": plazo})
        assert est == 200, (est, b)
        archivo = base64.b64decode(b["ArchivoBase64"], validate=True)
    chequear(len(archivo) == b["TamanoBytes"], f"{plazo}: TamanoBytes no coincide")
    ruta = os.path.join(SAL, f"sandbox-{plazo}.xlsx" if not CONTROL_NEGATIVO else f"control-negativo-{plazo}.xlsx"); open(ruta, "wb").write(archivo)
    # (1) contra el esperado local
    loc = zipfile.ZipFile(os.path.join(SAL, f"local-{plazo}.xlsx")); sbx = zipfile.ZipFile(io.BytesIO(archivo))
    nl, ns = set(loc.namelist()), set(sbx.namelist())
    chequear(nl == ns, f"{plazo}: partes distintas {nl ^ ns}")
    dif = [n for n in sorted(nl & ns) if loc.read(n) != sbx.read(n)]
    chequear(not dif, f"{plazo}: partes con bytes distintos {dif}")
    # (2) oraculo independiente
    wb = openpyxl.load_workbook(io.BytesIO(archivo))
    cu, simples = esperado(plazo)
    ok_simples = 0
    for (hoja, celda), esp in simples.items():
        v = wb[hoja][celda].value
        if chequear(igual(v, esp), f"{plazo}: {hoja}!{celda} = {v!r} (tipo {type(v).__name__}), esperado {esp!r}"):
            ok_simples += 1
    ws = wb[CONTROL]
    cuotas_ok = 0
    for i, f, m, fp, mp in cu:
        r = 6 + i
        fila = [ws.cell(r, c).value for c in range(1, 6)]   # D tiene formato de moneda en la plantilla (hallazgo 4 del spike 1): openpyxl lo lee como serial
        if chequear(fila[0] == i and fila[1] == f and Decimal(str(fila[2])) == m and
                    ((fp is None and fila[3] is None) or (fp is not None and fila[3] == (fp - dt.datetime(1899, 12, 30)).days)) and ((mp is None and fila[4] is None) or Decimal(str(fila[4])) == mp), f"{plazo}: cuota {i} = {fila}"):
            cuotas_ok += 1
    sobrantes_limpias = 0; estilo_ok = 0
    wt = wb_t[CONTROL]
    for r in range(7 + plazo, 67):
        vac = all(ws.cell(r, c).value is None and ws.cell(r, c).data_type != "f" for c in range(1, 6))
        est_igual = all(ws.cell(r, c)._style == wt.cell(r, c)._style or ws.cell(r, c).number_format == wt.cell(r, c).number_format and ws.cell(r, c).border == wt.cell(r, c).border for c in range(1, 6))
        sobrantes_limpias += vac; estilo_ok += est_igual
    n_sobr = 60 - plazo
    chequear(sobrantes_limpias == n_sobr, f"{plazo}: filas sobrantes limpias {sobrantes_limpias}/{n_sobr}")
    chequear(estilo_ok == n_sobr, f"{plazo}: filas sobrantes con estilo/bordes de la plantilla {estilo_ok}/{n_sobr}")
    for hoja in (FORMATO, CONTROL):
        chequear(sorted(map(str, wb[hoja].merged_cells.ranges)) == sorted(map(str, wb_t[hoja].merged_cells.ranges)), f"{plazo}: combinadas distintas en {hoja}")
        chequear(wb[hoja].print_area == wb_t[hoja].print_area, f"{plazo}: area de impresion distinta en {hoja}: {wb[hoja].print_area} vs {wb_t[hoja].print_area}")
        chequear(len(wb[hoja]._images) == len(wb_t[hoja]._images), f"{plazo}: imagenes distintas en {hoja}")
        chequear(wb[hoja].column_dimensions.keys() == wb_t[hoja].column_dimensions.keys() and
                 all(wb[hoja].column_dimensions[k].width == wb_t[hoja].column_dimensions[k].width for k in wb_t[hoja].column_dimensions), f"{plazo}: anchos de columna distintos en {hoja}")
    media_t = {n: plantilla_zip.read(n) for n in plantilla_zip.namelist() if n.startswith("xl/media/") or n.startswith("xl/drawings/")}
    chequear(all(sbx.read(n) == v for n, v in media_t.items()), f"{plazo}: imagen/dibujo distinto de la plantilla")
    xml_hojas = b"".join(sbx.read(n) for n in sbx.namelist() if n.startswith("xl/worksheets/sheet"))
    chequear(b"TODAY" not in xml_hojas.upper(), f"{plazo}: quedo TODAY() en alguna hoja")
    chequear("xl/calcChain.xml" not in sbx.namelist(), f"{plazo}: sigue calcChain.xml")
    formulas = sum(1 for h in wb for fila in h.iter_rows() for c in fila if c.data_type == "f")
    full_calc = b"fullCalcOnLoad" in sbx.read("xl/workbook.xml")
    # (3) LibreOffice -> PDF
    pdfdir = os.path.join(SAL, f"pdf-{plazo}"); os.makedirs(pdfdir, exist_ok=True)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", pdfdir, ruta], capture_output=True, timeout=180)
    pdf = os.path.join(pdfdir, os.path.basename(ruta).replace(".xlsx", ".pdf"))
    info = subprocess.run(["pdfinfo", pdf], capture_output=True, text=True).stdout
    paginas = int([l for l in info.splitlines() if l.startswith("Pages:")][0].split()[1])
    texto = subprocess.run(["pdftotext", "-layout", pdf, "-"], capture_output=True, text=True).stdout
    chequear("Ana Lucia Martinez Ruiz (ficticio)" in texto and "Oficial de Prueba" in texto, f"{plazo}: el PDF no trae el nombre/oficial")
    chequear(f"Plazo de {plazo} cuotas" in texto, f"{plazo}: el PDF no trae el comentario")
    resumen.append(f"plazo {plazo}: {len(archivo)} bytes | partes iguales al esperado local {len(nl & ns) - len(dif)}/{len(nl)} | celdas simples/totales ok {ok_simples}/{len(simples)} | cuotas ok {cuotas_ok}/{plazo} | "
                   f"filas sobrantes limpias {sobrantes_limpias}/{n_sobr} (estilo {estilo_ok}/{n_sobr}) | formulas restantes {formulas} fullCalcOnLoad={full_calc} | PDF {paginas} paginas")
print("\n".join(resumen))
print("FALLOS:" if fallos else "SIN FALLOS", *fallos, sep="\n  ")
sys.exit(1 if fallos else 0)
